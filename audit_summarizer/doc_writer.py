"""
doc_writer.py
----------------
Turns the LLM's markdown-ish summary text into a polished, committee-ready
Word document (with a cover page), and optionally converts that to PDF.

Handles right-to-left (Persian/Farsi) text correctly: paragraphs are marked
bidi=True and the paragraph alignment is set to RIGHT, which is what Word
needs to lay Persian text out properly.
"""

from __future__ import annotations

import datetime as _dt
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

RTL_FONT = "B Nazanin"       # common Persian font; falls back gracefully if absent
LATIN_FONT = "Calibri"
ACCENT_COLOR = RGBColor(0x1F, 0x3A, 0x5F)


def _set_rtl(paragraph, *, align: bool = True) -> None:
    """Mark a paragraph as right-to-left (needed for correct Persian rendering).

    Only forces RIGHT alignment when `align` is True; callers that already set
    a specific alignment (e.g. CENTER on the cover page) should pass align=False
    so this doesn't clobber it.
    """
    if align:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.makeelement(qn("w:bidi"), {})
    pPr.append(bidi)
    for run in paragraph.runs:
        rPr = run._r.get_or_add_rPr()
        rtl_el = rPr.makeelement(qn("w:rtl"), {})
        rPr.append(rtl_el)


def _style_run(run, *, size=11, bold=False, color=None, font=RTL_FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    # Ensure complex-script font is also set, or Word may use a Latin font for Persian glyphs.
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:cs"), font)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)


_INLINE_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _add_runs_with_inline_bold(paragraph, text: str, *, size=11, base_bold=False, color=None):
    """Split `text` on **bold** markers and add each segment as its own run."""
    pos = 0
    for match in _INLINE_BOLD_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            _style_run(run, size=size, bold=base_bold, color=color)
        run = paragraph.add_run(match.group(1))
        _style_run(run, size=size, bold=True, color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _style_run(run, size=size, bold=base_bold, color=color)


def _add_paragraph(doc: Document, text: str, *, size=11, bold=False, color=None, rtl=True):
    p = doc.add_paragraph()
    _add_runs_with_inline_bold(p, text, size=size, base_bold=bold, color=color)
    if rtl:
        _set_rtl(p)
    return p


def _add_bullet(doc: Document, text: str, *, rtl=True):
    p = doc.add_paragraph(style="List Bullet")
    _add_runs_with_inline_bold(p, text, size=11)
    if rtl:
        _set_rtl(p)
    return p


def _add_cover_page(doc: Document, *, title: str, organization: str | None, meeting_context: str | None):
    today = _dt.date.today().strftime("%Y-%m-%d")

    # Spacer to push the title roughly to the vertical center of the page.
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run(title)
    _style_run(run, size=26, bold=True, color=ACCENT_COLOR)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl(p, align=False)

    if organization:
        p2 = doc.add_paragraph()
        run2 = p2.add_run(organization)
        _style_run(run2, size=16, bold=False)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_rtl(p2, align=False)

    p3 = doc.add_paragraph()
    run3 = p3.add_run(f"تاریخ تهیه گزارش: {today}")
    _style_run(run3, size=12)
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl(p3, align=False)

    if meeting_context:
        p4 = doc.add_paragraph()
        run4 = p4.add_run(meeting_context)
        _style_run(run4, size=12)
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_rtl(p4, align=False)

    disclaimer = (
        "این سند خلاصه‌ای است که با کمک هوش مصنوعی از گزارش حسابرسی اصلی تهیه شده است "
        "و صرفاً برای تسهیل بررسی در جلسه است. سند اصلی حسابرسی مرجع نهایی و رسمی محسوب می‌شود."
    )
    for _ in range(6):
        doc.add_paragraph()
    p5 = doc.add_paragraph()
    run5 = p5.add_run(disclaimer)
    _style_run(run5, size=9, color=RGBColor(0x66, 0x66, 0x66))
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl(p5, align=False)

    doc.add_page_break()


_LIST_ITEM_RE = re.compile(r"^\s*(?:[-*•]|\d+[.\)])\s+(.*)")
_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)")


def render_summary_to_docx(
    summary_markdown: str,
    output_path: str | Path,
    *,
    title: str = "خلاصه گزارش حسابرسی برای جلسه کمیته",
    organization: str | None = None,
    meeting_context: str | None = None,
    source_filename: str | None = None,
) -> Path:
    """
    Convert the LLM's markdown-style summary into a formatted .docx file with
    a cover page. Recognises '#'/'##'/'###' headings and '-'/'*'/'1.' list items;
    everything else is rendered as a normal paragraph.
    """
    output_path = Path(output_path)
    doc = Document()

    # Base document defaults (helps when a heading line falls back to Normal style).
    normal = doc.styles["Normal"]
    normal.font.name = RTL_FONT
    normal.font.size = Pt(11)

    _add_cover_page(doc, title=title, organization=organization, meeting_context=meeting_context)

    if source_filename:
        _add_paragraph(
            doc, f"سند منبع: {source_filename}", size=9, color=RGBColor(0x88, 0x88, 0x88)
        )
        doc.add_paragraph()

    for raw_line in summary_markdown.splitlines():
        line = raw_line.strip()
        if not line:
            doc.add_paragraph()
            continue

        heading_match = _HEADING_RE.match(line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            size = {1: 18, 2: 15, 3: 13}.get(level, 12)
            p = _add_paragraph(doc, text, size=size, bold=True, color=ACCENT_COLOR)
            # small rule under top-level headings for visual separation
            if level <= 2:
                pPr = p._p.get_or_add_pPr()
                pBdr = pPr.makeelement(qn("w:pBdr"), {})
                bottom = pPr.makeelement(qn("w:bottom"), {
                    qn("w:val"): "single", qn("w:sz"): "6",
                    qn("w:space"): "4", qn("w:color"): "1F3A5F",
                })
                pBdr.append(bottom)
                pPr.append(pBdr)
            continue

        list_match = _LIST_ITEM_RE.match(line)
        if list_match:
            _add_bullet(doc, list_match.group(1).strip())
            continue

        _add_paragraph(doc, line)

    doc.save(str(output_path))
    return output_path


def convert_docx_to_pdf(docx_path: str | Path, output_dir: str | Path | None = None) -> Path:
    """Convert a .docx to .pdf using headless LibreOffice."""
    docx_path = Path(docx_path)
    output_dir = Path(output_dir) if output_dir else docx_path.parent

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(
            "برای تولید PDF، LibreOffice (soffice) لازم است اما روی سیستم یافت نشد."
        )

    cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    if result.returncode != 0:
        raise RuntimeError(f"تبدیل به PDF ناموفق بود: {result.stderr.strip()}")

    pdf_path = output_dir / (docx_path.stem + ".pdf")
    if not pdf_path.exists():
        raise RuntimeError("فایل PDF خروجی یافت نشد.")
    return pdf_path
