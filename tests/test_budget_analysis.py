"""تست‌های کارگاه «تحلیل بودجه» (خط سه‌مرحله‌ای + یکپارچگی با داشبورد/پروژه‌ها).

راهبرد تست، دقیقاً مثل سایر کارگاه‌های پروژه:

* مرحله‌ی ۲ (تنها بخشی که به مدل زبانی نیاز دارد) با یک تابع جعلی جایگزین می‌شود؛
  بنابراین تست‌ها نه کلید API می‌خواهند و نه شبکه.
* مراحل ۱ و ۳ (استخراج ساختارمند و رندر Word) **واقعی** اجرا می‌شوند تا دو مورد از
  «تعریف انجام‌شده» مستقیماً آزموده شود: کارکرد ورودی xlsx و ورودی PDF چاپ‌شده از
  Excel، و ظاهر شدن «در اسناد موجود نیست» به‌جای داده‌ی جعلی.
"""
from __future__ import annotations

import io
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Border, Side

from budget_analysis import MISSING_TEXT, NOT_COMPUTABLE_TEXT
from budget_analysis.config import BudgetConfig
from budget_analysis.extraction import extract_bundle, extract_document
from budget_analysis.prompt import build_messages
from budget_analysis.report import ReportContext, render_report_bytes
from budget_analysis.schemas import BudgetAnalysisReport

SOFFICE_CANDIDATES = (
    r"C:/Program Files/LibreOffice/program/soffice.exe",
    "soffice",
    "libreoffice",
)


# ---------------------------------------------------------------------------
# ساخت اسناد نمونه (واقع‌گرایانه: عنوان بخش ادغام‌شده، سرصفحه‌ی چندستونی، واحد)
# ---------------------------------------------------------------------------
def _form_1_sheet(ws, current_year: int, base_year: int, factor: float) -> None:
    ws.title = "فرم 1"
    ws.append(["فرم 1 - منابع و مصارف (میلیون ریال)"])
    ws.append(["ردیف", "عنوان", f"اصلاحیه بودجه سال {current_year}", f"عملکرد سال {base_year}"])
    rows = [
        (1, "منابع عمومی", 100_000, 90_000),
        (2, "درآمد اختصاصی", 20_000, 18_000),
        (3, "جمع منابع", 120_000, 108_000),
        (4, "اعتبارات هزینه ای", 84_000, 76_000),
        (5, "تملک دارایی های سرمایه ای", 36_000, 32_000),
        (6, "جمع مصارف", 120_000, 108_000),
    ]
    for number, title, current, previous in rows:
        ws.append([number, title, int(current * factor), previous])


def _form_5_1_sheet(ws, current_year: int, base_year: int, factor: float) -> None:
    ws.title = "فرم 5-1"
    ws.append(["فرم 5-1 - مصارف و تفکیک هزینه ها (میلیون ریال)"])
    ws.append(["ردیف", "عنوان قلم", f"اصلاحیه بودجه سال {current_year}", f"عملکرد سال {base_year}"])
    ws.append(["هزینه های پرسنلی"])
    rows = [
        (16, "حقوق و دستمزد", 12_500, 9_800),
        (17, "مزایای غیرمستمر", 3_200, 2_100),
        (19, "هزینه های رفاهی", 900, 700),
        (30, "بازخرید خدمت", 500, 480),
    ]
    for number, title, current, previous in rows:
        ws.append([number, title, int(current * factor), previous])
    ws.append([31, "جمع هزینه های پرسنلی", int(17_100 * factor), 13_080])


def _form_6_sheet(ws, current_year: int, base_year: int, factor: float) -> None:
    ws.title = "فرم 6"
    ws.append(["فرم 6 - آمار کارکنان و اعضای هیات علمی"])
    ws.append(["شرح", f"سال {current_year}", f"سال {base_year}"])
    ws.append(["کارکنان غیر هیات علمی", 120, 115])
    ws.append(["اعضای هیات علمی", 18, 17])
    _ = factor


def _border_all(ws, max_row: int, max_col: int) -> None:
    thin = Side(style="thin")
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def make_budget_workbook(
    path: Path,
    *,
    current_year: int,
    base_year: int,
    factor: float = 1.0,
    include_form_6: bool = False,
) -> Path:
    """یک سند اصلاحیه بودجه‌ی نمونه با فرم‌های ۱ و ۵-۱ (و به‌صورت اختیاری فرم ۶)."""
    workbook = Workbook()
    _form_1_sheet(workbook.active, current_year, base_year, factor)
    _form_5_1_sheet(workbook.create_sheet(), current_year, base_year, factor)
    if include_form_6:
        _form_6_sheet(workbook.create_sheet(), current_year, base_year, factor)
    for ws in workbook.worksheets:
        _border_all(ws, ws.max_row, ws.max_column)
    workbook.save(path)
    return path


def _soffice() -> str | None:
    for candidate in SOFFICE_CANDIDATES:
        resolved = shutil.which(candidate)
        if resolved:
            return resolved
    return None


def convert_to_pdf(source: Path, out_dir: Path) -> Path | None:
    """xlsx → PDF با LibreOffice (شبیه‌سازی «PDF چاپ‌شده از اکسل»)."""
    soffice = _soffice()
    if soffice is None:
        return None
    try:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(source)],
            capture_output=True,
            text=True,
            timeout=180,
        )
    except (OSError, subprocess.TimeoutExpired):
        return None
    if result.returncode != 0:
        return None
    pdf_path = out_dir / (source.stem + ".pdf")
    return pdf_path if pdf_path.exists() else None


# ---------------------------------------------------------------------------
# مرحله‌ی ۲ جعلی
# ---------------------------------------------------------------------------
def fake_report_payload() -> dict:
    """خروجی sample مرحله‌ی ۲ با هر دو حالت «انحراف» و «فاقد داده کافی»."""
    return {
        "meta": {
            "organization": "پارک علم و فناوری تست",
            "current_year": "1405",
            "base_year": "1404",
            "source_documents": [],
        },
        "executive_summary": {
            "overall_status": "وضعیت کلی نیازمند بررسی مدیریتی است.",
            "top_deviation": "رشد ۸۴٫۹ درصدی قلم حقوق و دستمزد در فرم ۵-۱",
            "top_financial_risk": "فشار هزینه‌های پرسنلی بر منابع قابل تخصیص به فناوری",
            "top_structural_risk": "ساختار هزینه‌ای متکی بر اقلام پرسنلی",
            "top_mission_issue": "کاهش سهم اعتبارات حمایت مستقیم از واحدهای فناور",
            "top_recommended_decision": "تعیین سقف رشد اقلام پرسنلی در جلسه",
        },
        "axis_dashboard": [
            {
                "axis": "رشد حقوق، مزایا و اقلام مرتبط",
                "status": "هشدار مدیریتی",
                "importance": "بااهمیت",
                "significant_findings_count": 1,
                "management_summary": "رشد ۸۴٫۹ درصدی ردیف حقوق و دستمزد، ۴۴٫۹ واحد درصد بالاتر از سقف ۴۰٪",
            }
        ],
        "error_matrix": [
            {
                "criterion_id": "C3.1",
                "axis": "رشد حقوق، مزایا و اقلام مرتبط",
                "criterion": "رشد اقلام پرسنلی نسبت به عملکرد سال پایه",
                "result": "رشد ۸۴٫۹ درصد در برابر حد بالای ۴۰ درصد",
                "status": "هشدار مدیریتی",
                "importance": "بااهمیت",
                "form": "فرم 5-1",
                "section": "هزینه های پرسنلی",
                "row": "16 حقوق و دستمزد",
                "column": "اصلاحیه بودجه سال 1405",
                "base_value": 9800,
                "current_value": 18125,
                "threshold": 13720,
                "absolute_deviation": 4405,
                "percentage_deviation": 44.9,
                "comparison_type": "بودجه جاری در مقابل عملکرد سال پایه",
                "base_document": "base.xlsx",
                "comparison_document": "current.xlsx",
                "base_year": "1404",
                "comparison_year": "1405",
                "evidence": "ردیف ۱۶ فرم ۵-۱: ۹٫۸۰۰ به ۱۸٫۱۲۵ میلیون ریال",
                "action": "ارائه مستندات افزایش و تعیین سقف مصوب",
            },
            {
                "criterion_id": "C4.5",
                "axis": "ساختار هزینه‌های نیروی انسانی",
                "criterion": "هزینه سرانه نیروی انسانی",
                "result": "",
                "status": "فاقد داده کافی",
                "importance": "عادی",
                "form": "",
                "row": "",
                "column": "",
                "evidence": "تعداد کارکنان در سند سال جاری موجود نیست",
            },
        ],
        "top_findings": [
            {
                "rank": 1,
                "axis": "رشد حقوق، مزایا و اقلام مرتبط",
                "criterion": "رشد اقلام پرسنلی",
                "location": "فرم 5-1 ← ردیف 16",
                "observed_value": 84.9,
                "reference_value": 40,
                "deviation": 44.9,
                "deviation_unit": "واحد درصد",
                "importance": "بااهمیت",
                "management_message": "عبور قطعی از سقف مجاز رشد اقلام پرسنلی",
            }
        ],
        "significant_findings": [
            {
                "number": 1,
                "title": "رشد غیرعادی اقلام پرسنلی",
                "subject": "رشد حقوق و دستمزد از ۹٫۸۰۰ به ۱۸٫۱۲۵ میلیون ریال",
                "evidence": {
                    "form": "فرم 5-1",
                    "section": "هزینه های پرسنلی",
                    "row": "16 حقوق و دستمزد",
                    "column": "اصلاحیه بودجه سال 1405",
                    "base_value": 9800,
                    "current_value": 18125,
                    "threshold": 13720,
                    "absolute_deviation": 4405,
                    "percentage_deviation": 44.9,
                },
                "assessment": "خارج از دامنه‌ی مجاز پارامتر رشد",
                "management_importance": "فشار بر منابع قابل تخصیص به فعالیت‌های فناورانه",
                "risk": "کاهش فضای مالی برنامه‌های توسعه فناوری",
                "action": "اصلاح تخصیص یا ارائه مستند تغییر حجم نیروی انسانی",
            }
        ],
        "risks": [
            {
                "description": "فشار هزینه‌های پرسنلی بر منابع فناوری",
                "approximate_amount": "قابل تعیین از اسناد موجود نیست",
                "importance": "بااهمیت",
                "nature": "مالی",
            }
        ],
        "items_needing_decision": [
            {
                "subject": "سقف رشد اقلام پرسنلی",
                "question": "آیا رشد ۸۴٫۹ درصدی ردیف حقوق و دستمزد در جلسه تأیید می‌شود؟",
                "proposed_action": "تعیین سقف مصوب رشد و الزام به ارائه مستندات",
            }
        ],
        "items_needing_clarification": [
            {
                "subject": "آمار کارکنان سال جاری",
                "available_data": "آمار کارکنان فقط در سند سال پایه موجود است",
                "missing_or_conflicting_data": "تعداد کارکنان غیر هیأت علمی سال جاری",
                "location": "فرم 6",
                "reason": "محاسبه‌ی هزینه سرانه نیروی انسانی بدون این داده ممکن نیست",
                "required_document": "فرم ۶ آمار کارکنان سال جاری",
            }
        ],
        "closing_notes": ["پیشنهاد می‌شود مستندات تغییر حجم نیروی انسانی پیش از جلسه ارائه شود."],
    }


def _fake_analyze(report_payload: dict | None = None):
    """جایگزین مرحله‌ی ۲: خروجی نمونه را اعتبارسنجی و برمی‌گرداند."""
    payload = report_payload or fake_report_payload()

    def _analyze(bundle, config=None, *, organization=None, meeting_context=None, llm=None, report_stage=None, **kwargs):
        if report_stage is not None:
            report_stage("مرحله ۲ از ۳: تحلیل معیارها با هوش مصنوعی...")
        report = BudgetAnalysisReport.model_validate(payload)
        if organization:
            report.meta.organization = organization
        if bundle.base_year:
            report.meta.base_year = bundle.base_year
        if bundle.current_year:
            report.meta.current_year = bundle.current_year
        return report

    return _analyze


def _patch_sync_start(monkeypatch, *, analyze=None, error: str | None = None):
    """``start_budget_job`` را با نسخه‌ی همزمان جایگزین می‌کند (بدون ترد، بدون LLM)."""
    import budget_analysis.pipeline as budget_pipeline

    def _fake_start(job, file_paths, workdir, **kwargs):
        if error:
            job["status"] = "error"
            job["error"] = error
            budget_pipeline.cleanup_workdir(workdir)
            return
        try:
            job["result"] = budget_pipeline.run_budget_analysis(
                file_paths,
                workdir=workdir,
                config=kwargs.get("config"),
                organization=kwargs.get("organization"),
                meeting_context=kwargs.get("meeting_context"),
                report_date=kwargs.get("report_date"),
                job=job,
                analyze=analyze or _fake_analyze(),
                filenames=kwargs.get("filenames"),
            )
            job["status"] = "done"
        except Exception as exc:  # noqa: BLE001
            job["status"] = "error"
            job["error"] = str(exc)
        finally:
            budget_pipeline.cleanup_workdir(workdir)

    monkeypatch.setattr(budget_pipeline, "start_budget_job", _fake_start)


def _xlsx_upload(name: str, path: Path):
    return (name, (path.name, io.BytesIO(path.read_bytes()), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))


def _pdf_upload(name: str, path: Path):
    return (name, (path.name, io.BytesIO(path.read_bytes()), "application/pdf"))


# ---------------------------------------------------------------------------
# مرحله ۱: استخراج
# ---------------------------------------------------------------------------
def test_extraction_reads_forms_rows_and_cell_locations(tmp_path):
    base = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    current = make_budget_workbook(tmp_path / "current.xlsx", current_year=1405, base_year=1404, factor=1.45)

    bundle = extract_bundle(
        {"base_year": base, "current_year": current},
        config=BudgetConfig(),
        filenames={"base_year": "base.xlsx", "current_year": "current.xlsx"},
    )

    assert bundle.base_year == "1404"
    assert bundle.current_year == "1405"

    current_document = bundle.document("current_year")
    assert current_document is not None
    assert current_document.unit == "میلیون ریال"

    form_keys = {form.form_key for form in current_document.forms}
    assert {"form_1", "form_5_1"} <= form_keys

    salaries = [
        cell
        for cell in current_document.cells
        if cell.form_key == "form_5_1" and cell.row_number == 16
    ]
    assert salaries, "ردیف ۱۶ فرم ۵-۱ باید استخراج شود"
    values = {cell.year: cell.value for cell in salaries}
    assert values["1405"] == pytest.approx(12_500 * 1.45)
    assert values["1404"] == pytest.approx(9_800)
    # محل داده (ردیابی) ثبت شده است
    assert all(cell.cell_ref for cell in salaries)
    assert {cell.section for cell in salaries} == {"هزینه های پرسنلی"}


def test_extraction_marks_missing_forms_without_fabricating_data(tmp_path):
    """فرم‌هایی که در سند نیستند فقط «شناسایی‌نشده» علامت می‌خورند و داده‌ای ساخته نمی‌شود."""
    base = make_budget_workbook(
        tmp_path / "base.xlsx", current_year=1404, base_year=1403, include_form_6=True
    )
    current = make_budget_workbook(tmp_path / "current.xlsx", current_year=1405, base_year=1404)

    bundle = extract_bundle({"base_year": base, "current_year": current}, config=BudgetConfig())
    base_document = bundle.document("base_year")
    current_document = bundle.document("current_year")

    assert "form_6" in {form.form_key for form in base_document.forms}
    assert "form_6" in current_document.missing_form_keys
    # هیچ سلولی برای فرم شناسایی‌نشده ساخته نمی‌شود
    assert all(cell.form_key != "form_6" for cell in current_document.cells)
    # و مقدار صفر جعلی هم ساخته نمی‌شود
    assert all(cell.value != 0 or cell.raw_value for cell in current_document.cells)


# ---------------------------------------------------------------------------
# مرحله ۱: ورودی PDF چاپ‌شده از اکسل
# ---------------------------------------------------------------------------
def test_extraction_accepts_pdf_of_excel(tmp_path):
    xlsx = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    pdf = convert_to_pdf(xlsx, tmp_path)
    if pdf is None:
        pytest.skip("LibreOffice برای تولید PDF آزمون در این محیط در دسترس نیست.")

    document = extract_document(pdf, "base_year")

    assert document.extraction_method.startswith("pdf")
    # نام شیت/فرم باید به‌صورت معنایی شناسایی شود (نه با موقعیت سلول)
    assert {"form_1", "form_5_1"} <= {form.form_key for form in document.forms}

    values = [cell.value for cell in document.cells]
    # ارقام باید سالم بازیابی شوند (ستون‌های «اصلاحیه» و «عملکرد»)
    assert 100_000 in values or 90_000 in values
    # و ردیف‌های پرسنلی قابل شناسایی باشند
    salary_rows = [cell for cell in document.cells if cell.form_key == "form_5_1"]
    assert salary_rows
    assert any(cell.value in {12_500, 9_800} for cell in salary_rows)


# ---------------------------------------------------------------------------
# مرحله ۲: اعتبارسنجی سخت‌گیرانه و یک تلاش مجدد
# ---------------------------------------------------------------------------
class _ScriptedLLM:
    """کلاینت جعلی که پاسخ‌های از پیش تعیین‌شده برمی‌گرداند و پیام‌ها را ثبت می‌کند."""

    def __init__(self, responses: list[object]) -> None:
        self.responses = list(responses)
        self.calls: list[list[dict[str, str]]] = []

    def complete_json(self, messages):
        self.calls.append([dict(message) for message in messages])
        if not self.responses:
            raise AssertionError("پاسخ جعلی بیشتری وجود ندارد")
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response


def _bundle(tmp_path):
    base = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    current = make_budget_workbook(tmp_path / "current.xlsx", current_year=1405, base_year=1404)
    return extract_bundle({"base_year": base, "current_year": current}, config=BudgetConfig())


def test_analysis_retries_once_with_validation_error(tmp_path):
    from budget_analysis.analysis import analyze_budget

    invalid = dict(fake_report_payload())
    invalid["error_matrix"] = [
        {
            "criterion_id": "C3.1",
            "axis": "رشد حقوق",
            "result": "انحراف",
            "status": "هشدار مدیریتی",
            "importance": "بااهمیت",
            # فرم و ردیف خالی است => ردیف دارای انحراف بدون محل داده
        }
    ]
    llm = _ScriptedLLM([invalid, fake_report_payload()])

    report = analyze_budget(_bundle(tmp_path), BudgetConfig(), llm=llm)

    assert len(llm.calls) == 2, "خروجی نامعتبر باید یک‌بار تکرار شود"
    assert "اصلاح خروجی قبلی" in llm.calls[1][1]["content"]
    assert "ردیف دارای انحراف" in llm.calls[1][1]["content"]
    assert len(report.error_matrix) == 2


def test_analysis_fails_with_persian_error_when_output_stays_invalid(tmp_path):
    from budget_analysis.analysis import BudgetAnalysisError, analyze_budget

    invalid = dict(fake_report_payload())
    invalid["error_matrix"] = [
        {"criterion_id": "C1.1", "axis": "محور", "result": "انحراف", "status": "مغایرت بااهمیت", "importance": "بااهمیت"}
    ]
    llm = _ScriptedLLM([invalid, invalid])

    with pytest.raises(BudgetAnalysisError) as excinfo:
        analyze_budget(_bundle(tmp_path), BudgetConfig(), llm=llm)

    assert "دو تلاش" in str(excinfo.value)
    assert len(llm.calls) == 2


def test_analysis_rejects_status_outside_vocabulary(tmp_path):
    from budget_analysis.analysis import BudgetAnalysisError, analyze_budget

    invalid = dict(fake_report_payload())
    invalid["axis_dashboard"] = [
        {"axis": "محور", "status": "خطرناک", "importance": "بااهمیت", "significant_findings_count": 1}
    ]
    llm = _ScriptedLLM([invalid, invalid])

    with pytest.raises(BudgetAnalysisError):
        analyze_budget(_bundle(tmp_path), BudgetConfig(), llm=llm)


def test_service_error_is_not_retried_and_surfaces_persian_message(tmp_path):
    """خطای سرویس هوش مصنوعی با تکرار ساختاری بهتر نمی‌شود؛ همان پیام با یک تلاش برمی‌گردد."""
    from budget_analysis.analysis import BudgetAnalysisError, analyze_budget
    from budget_analysis.llm import BudgetLLMError

    llm = _ScriptedLLM([BudgetLLMError("سهمیه سرویس هوش مصنوعی تمام شده است.")])

    with pytest.raises(BudgetAnalysisError) as excinfo:
        analyze_budget(_bundle(tmp_path), BudgetConfig(), llm=llm)

    assert "سهمیه" in str(excinfo.value)
    assert len(llm.calls) == 1, "خطای سرویس نباید تلاش مجدد ساختاری بگیرد"


def test_missing_api_key_produces_clear_error(monkeypatch):
    from budget_analysis.llm import BudgetAnalysisLLM, BudgetLLMError, BudgetLLMSettings

    for variable in ("AUDIT_REPORT_LLM_API_KEY", "API_KEY_OPENROUTER", "B_AI_API_KEY"):
        monkeypatch.delenv(variable, raising=False)

    with pytest.raises(BudgetLLMError) as excinfo:
        BudgetAnalysisLLM(BudgetLLMSettings()).complete_json([{"role": "user", "content": "x"}])

    assert "کلید" in str(excinfo.value)


def test_job_lifecycle_removes_temp_workdir_and_intermediate_files(tmp_path):
    """JSONهای میانی و فایل‌های ورودی پس از پایان کار همراه پوشهٔ موقت حذف می‌شوند."""
    import time

    import budget_analysis.pipeline as budget_pipeline

    base = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    current = make_budget_workbook(tmp_path / "current.xlsx", current_year=1405, base_year=1404)

    workdir = budget_pipeline.make_temp_workdir()
    paths = {"base_year": workdir / "base.xlsx", "current_year": workdir / "current.xlsx"}
    paths["base_year"].write_bytes(base.read_bytes())
    paths["current_year"].write_bytes(current.read_bytes())

    job = budget_pipeline.new_budget_job()
    budget_pipeline.start_budget_job(
        job,
        paths,
        workdir,
        config=BudgetConfig(),
        organization="پارک تست",
        analyze=_fake_analyze(),
        filenames={"base_year": "base.xlsx", "current_year": "current.xlsx"},
    )

    deadline = time.time() + 30
    while job.get("status") not in ("done", "error") and time.time() < deadline:
        time.sleep(0.05)

    assert job["status"] == "done", job.get("error")
    assert job["result"]["docx_bytes"][:2] == b"PK"
    assert not workdir.exists(), "پوشهٔ کار موقت باید پس از پایان پردازش حذف شود"


def test_progress_reflects_the_three_stages():
    from api.services import budget_service

    def progress(stage: str) -> int:
        return budget_service.estimate_progress({"status": "running", "stage": stage, "logs": []})

    assert progress("در حال شروع...") == 3
    assert progress("مرحله ۱ از ۳: استخراج ساختارمند فرم‌ها...") == 12
    assert progress("مرحله ۲ از ۳: تحلیل معیارها...") == 55
    assert progress("مرحله ۳ از ۳: تولید گزارش مدیریتی Word...") == 88
    assert progress("پردازش با موفقیت به اتمام رسید.") == 100


def test_missing_data_row_gets_standard_wording(tmp_path):
    """ردیف «فاقد داده کافی» با عبارت استاندارد پرامپت مرجع پر می‌شود."""
    from budget_analysis.schemas import BudgetAnalysisReport

    payload = fake_report_payload()
    payload["error_matrix"][1]["result"] = ""
    report = BudgetAnalysisReport.model_validate(payload)
    row = report.error_matrix[1]
    assert row.result == NOT_COMPUTABLE_TEXT
    assert row.form == MISSING_TEXT


# ---------------------------------------------------------------------------
# مرحله ۲: پرامپت
# ---------------------------------------------------------------------------
def test_prompt_carries_substantive_rules_and_thresholds(tmp_path):
    bundle = _bundle(tmp_path)
    messages = build_messages(bundle, BudgetConfig(), organization="پارک تست")
    system = messages[0]["content"]
    user = messages[1]["content"]

    # اصول ماهوی پرامپت مرجع باید در پرامپت باشند
    assert MISSING_TEXT in system and NOT_COMPUTABLE_TEXT in system
    assert "عدم حدس و جعل" in system
    assert "ردیابی کامل" in system
    assert "تحلیل سه‌لایه" in system
    assert "برخی" in system and "کلی‌گویی" in system
    assert "مغایرت بااهمیت" in system and "ریسک بااهمیت" in system
    assert "بسیار بالا" in system
    # فهرست محورها داده‌محور است، نه شش محور ثابت در متن پرامپت
    assert "سقف عملکرد نسبت به اعتبار" in system
    assert "ثبات تشکیلات و نیروی انسانی" in system
    assert "قابل توسعه است" in system
    # آستانه‌ها از پیکربندی می‌آیند
    assert "salary_upper_bound" in system
    # داده‌ی استخراج‌شده در پیام کاربر است
    assert "فرم 5-1" in user
    assert "پارک تست" in user


# ---------------------------------------------------------------------------
# مرحله ۳: رندر Word
# ---------------------------------------------------------------------------
def test_report_renders_required_sections_and_missing_data_text():
    import docx

    report = BudgetAnalysisReport.model_validate(fake_report_payload())
    data = render_report_bytes(
        report,
        ReportContext(
            organization="پارک علم و فناوری تست",
            base_year="1404",
            current_year="1405",
            unit="میلیون ریال",
            source_documents=["base.xlsx", "current.xlsx"],
            report_date="1405/06/26",
        ),
    )
    document = docx.Document(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    text = "\n".join(paragraphs)

    expected_order = [
        "توضیح درباره ماهیت گزارش",
        "اسناد مبنای تحلیل",
        "خلاصه مدیریتی",
        "جمع‌بندی وضعیت",
        "وضعیت محورهای پایش",
        "جزئیات نتایج محورها و موارد انحراف",
        "مهم‌ترین انحرافات عددی",
        "یافته‌های بااهمیت",
        "ریسک‌ها و آثار مالی",
        "موارد نیازمند بررسی و تصمیم‌گیری",
        "موارد نیازمند شفاف‌سازی",
        "نکات تکمیلی برای جلسه",
    ]
    positions = [text.index(heading) for heading in expected_order]
    assert positions == sorted(positions), "ترتیب بخش‌های گزارش مطابق قالب مصوب نیست"

    # ارقام فارسی و واحد در سربرگ جدول‌ها
    assert "۱۴۰۵" in text
    detail_table = document.tables[1]
    headers = [cell.text for cell in detail_table.rows[0].cells]
    assert any("میلیون ریال" in header for header in headers)

    # داده‌ی ناموجود به‌صورت «در اسناد موجود نیست» می‌آید، نه صفر یا خط تیره
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    body_text = text + "\n" + table_text
    assert MISSING_TEXT in body_text
    assert NOT_COMPUTABLE_TEXT in body_text


def test_report_docx_is_rtl_with_repeating_table_header():
    import zipfile

    report = BudgetAnalysisReport.model_validate(fake_report_payload())
    data = render_report_bytes(report, ReportContext(organization="پارک تست", unit="میلیون ریال"))
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        styles = archive.read("word/styles.xml").decode("utf-8")

    assert "<w:bidi" in xml, "پاراگراف‌های گزارش باید راست‌به‌چپ باشند"
    assert "bidiVisual" in xml, "جدول‌ها باید راست‌به‌چپ باشند"
    assert "tblHeader" in xml, "سربرگ جدول باید در صفحات بعدی تکرار شود"
    assert "fa-IR" in xml
    assert "Vazirmatn" in styles


# ---------------------------------------------------------------------------
# یکپارچگی با داشبورد/پروژه‌ها
# ---------------------------------------------------------------------------
def test_workshop_is_registered_and_listed_on_project_page(auth_client, project):
    from api.workshops import WORKSHOPS

    assert "budget-analysis" in WORKSHOPS
    page = auth_client.get(f"/projects/{project.id}")
    assert page.status_code == 200
    assert "تحلیل بودجه" in page.text
    assert f"/projects/{project.id}/budget-analysis" in page.text

    workshop_page = auth_client.get(f"/projects/{project.id}/budget-analysis")
    assert workshop_page.status_code == 200
    assert "اصلاحیه بودجه تفصیلی سال پایه" in workshop_page.text
    assert "اصلاحیه بودجه تفصیلی سال جاری" in workshop_page.text
    # اسکریپت صفحه هم باید در دسترس باشد (قالب آن را از مسیر استاتیک بار می‌کند)
    assert auth_client.get("/static/js/budget_analysis.js").status_code == 200


def test_budget_run_writes_history_and_downloadable_report(
    auth_client, db_session, project, monkeypatch, wait_for_run, tmp_path
):
    base = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    current = make_budget_workbook(tmp_path / "current.xlsx", current_year=1405, base_year=1404, factor=1.45)
    _patch_sync_start(monkeypatch)

    response = auth_client.post(
        f"/api/projects/{project.id}/budget-analysis/jobs",
        data={"organization": "پارک علم و فناوری تست"},
        files=[
            _xlsx_upload("budget_base_year", base),
            _xlsx_upload("budget_current_year", current),
        ],
    )
    assert response.status_code == 200, response.text
    job_id = response.json()["job_id"]
    run_id = response.json()["run_id"]

    status = auth_client.get(
        f"/api/projects/{project.id}/budget-analysis/jobs/{job_id}"
    ).json()
    assert status["status"] == "done", status
    assert status["report_ready"] is True
    assert status["progress"] == 100
    assert status["run_id"] == run_id

    results = auth_client.get(
        f"/api/projects/{project.id}/budget-analysis/jobs/{job_id}/results"
    ).json()
    assert results["matrix_total"] == 2
    assert results["deviation_total"] == 1
    assert results["axis_dashboard"][0]["axis"] == "رشد حقوق، مزایا و اقلام مرتبط"
    assert results["source_documents"] == ["base.xlsx", "current.xlsx"]
    assert results["unit"] == "میلیون ریال"

    # تاریخچه: ردیف تمام‌شده + فایل نتیجه روی دیسک
    job = wait_for_run(project.id, run_id)
    assert job["status"] == "done"
    assert job["download_ready"] is True
    assert job["result_summary"]["missing_data_count"] == 1
    assert job["result_summary"]["deviation_count"] == 1

    download = auth_client.get(
        f"/api/projects/{project.id}/budget-analysis/runs/{run_id}/download"
    )
    assert download.status_code == 200
    assert download.content[:2] == b"PK"  # یک docx واقعی (zip)
    assert "attachment" in download.headers["content-disposition"]


def test_budget_history_row_never_stores_input_files(
    auth_client, db_session, project, monkeypatch, wait_for_run, tmp_path
):
    base = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    current = make_budget_workbook(tmp_path / "current.xlsx", current_year=1405, base_year=1404)
    _patch_sync_start(monkeypatch)

    response = auth_client.post(
        f"/api/projects/{project.id}/budget-analysis/jobs",
        data={},
        files=[
            _xlsx_upload("budget_base_year", base),
            _xlsx_upload("budget_current_year", current),
        ],
    )
    run_id = response.json()["run_id"]
    wait_for_run(project.id, run_id)

    from api.db.models import WorkshopRun
    from api import storage

    run = db_session.get(WorkshopRun, run_id)
    db_session.refresh(run)
    stored = json.dumps(run.result_summary, ensure_ascii=False)

    # فقط نتیجه ذخیره می‌شود: مسیر/محتوای فایل ورودی و JSONهای میانی نباید بماند
    assert "psa_budget_" not in stored
    assert "01_extraction" not in stored and "02_analysis" not in stored
    assert run.result_file_path.startswith(f"project-{project.id}/")
    assert storage.resolve(run.result_file_path).name.endswith(".docx")
    # نام اسناد به‌عنوان متادیتای اجرا مجاز است، اما محتوای آن‌ها نه
    assert run.result_summary["source_filenames"] == ["base.xlsx", "current.xlsx"]


def test_budget_validation_requires_both_documents(auth_client, project, tmp_path):
    base = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    response = auth_client.post(
        f"/api/projects/{project.id}/budget-analysis/jobs",
        data={},
        files=[_xlsx_upload("budget_base_year", base)],
    )
    assert response.status_code == 400
    assert "الزامی" in response.json()["detail"]
    assert "سال جاری" in response.json()["detail"]


def test_budget_invalid_extension_returns_400(auth_client, project, tmp_path):
    base = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    response = auth_client.post(
        f"/api/projects/{project.id}/budget-analysis/jobs",
        data={},
        files=[
            _xlsx_upload("budget_base_year", base),
            ("budget_current_year", ("budget.txt", io.BytesIO(b"hello"), "text/plain")),
        ],
    )
    assert response.status_code == 400
    assert "فرمت" in response.json()["detail"]


def test_budget_job_is_scoped_to_its_project(auth_client, db_session, project, monkeypatch, tmp_path):
    base = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    current = make_budget_workbook(tmp_path / "current.xlsx", current_year=1405, base_year=1404)
    _patch_sync_start(monkeypatch)

    response = auth_client.post(
        f"/api/projects/{project.id}/budget-analysis/jobs",
        data={},
        files=[
            _xlsx_upload("budget_base_year", base),
            _xlsx_upload("budget_current_year", current),
        ],
    )
    job_id = response.json()["job_id"]

    from api.repositories import projects as projects_repo

    other = projects_repo.create(db_session, project.user_id, "پروژهٔ دوم")
    assert (
        auth_client.get(f"/api/projects/{other.id}/budget-analysis/jobs/{job_id}").status_code == 404
    )


def test_budget_error_is_recorded_in_history(auth_client, db_session, project, monkeypatch, wait_for_run, tmp_path):
    base = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    current = make_budget_workbook(tmp_path / "current.xlsx", current_year=1405, base_year=1404)
    _patch_sync_start(monkeypatch, error="سرویس هوش مصنوعی در دسترس نیست")

    response = auth_client.post(
        f"/api/projects/{project.id}/budget-analysis/jobs",
        data={},
        files=[
            _xlsx_upload("budget_base_year", base),
            _xlsx_upload("budget_current_year", current),
        ],
    )
    run_id = response.json()["run_id"]
    job = wait_for_run(project.id, run_id)
    assert job["status"] == "error"
    assert "هوش مصنوعی" in job["error"]

    history = auth_client.get(f"/projects/{project.id}/history")
    assert "سرویس هوش مصنوعی در دسترس نیست" in history.text


def test_budget_full_wiring_runs_through_real_job_runner(
    auth_client, db_session, project, monkeypatch, wait_for_run, tmp_path
):
    """مسیر واقعی router → service → ترد پس‌زمینه (بدون monkeypatch خود start_budget_job).

    این تست همان دسته خطایی را می‌گیرد که تست‌های جایگزین‌شده نمی‌بینند (مثلاً
    ناهم‌خوانی امضای توابع بین لایه‌ها).
    """
    import budget_analysis.pipeline as budget_pipeline

    base = make_budget_workbook(tmp_path / "base.xlsx", current_year=1404, base_year=1403)
    current = make_budget_workbook(tmp_path / "current.xlsx", current_year=1405, base_year=1404)

    # فقط مرحله‌ی مدل زبانی جعلی می‌شود؛ ترد، استخراج، رندر و ثبت تاریخچه واقعی‌اند.
    monkeypatch.setattr(budget_pipeline, "analyze_budget", _fake_analyze())
    workdirs_before = {str(path) for path in Path(tempfile.gettempdir()).glob(f"{budget_pipeline.WORKDIR_PREFIX}*")}

    response = auth_client.post(
        f"/api/projects/{project.id}/budget-analysis/jobs",
        data={"organization": "پارک علم و فناوری تست", "base_year": "1404", "current_year": "1405"},
        files=[
            _xlsx_upload("budget_base_year", base),
            _xlsx_upload("budget_current_year", current),
        ],
    )
    assert response.status_code == 200, response.text
    job_id = response.json()["job_id"]
    run_id = response.json()["run_id"]

    run = wait_for_run(project.id, run_id)
    assert run["status"] == "done", run

    status = auth_client.get(
        f"/api/projects/{project.id}/budget-analysis/jobs/{job_id}"
    ).json()
    assert status["status"] == "done"
    assert status["report_ready"] is True

    results = auth_client.get(
        f"/api/projects/{project.id}/budget-analysis/jobs/{job_id}/results"
    ).json()
    assert results["organization"] == "پارک علم و فناوری تست"
    assert results["base_year"] == "1404"
    assert results["current_year"] == "1405"
    assert results["matrix_total"] == 2
    assert run["result_file_path"]

    # هیچ پوشه‌ی کار موقتی پس از پایان اجرا باقی نمی‌ماند (ورودی‌ها و JSONهای میانی حذف شده‌اند)
    leftovers = {
        str(path) for path in Path(tempfile.gettempdir()).glob(f"{budget_pipeline.WORKDIR_PREFIX}*")
    } - workdirs_before
    assert not leftovers

    download = auth_client.get(
        f"/api/projects/{project.id}/budget-analysis/runs/{run_id}/download"
    )
    assert download.status_code == 200
    assert download.content[:2] == b"PK"
