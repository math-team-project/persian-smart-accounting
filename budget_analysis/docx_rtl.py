"""ابزارهای سطح-پایین ساخت سند Word فارسی راست‌به‌چپ (مرحله‌ی ۳).

این ماژول عمداً منطق RTL/فونت را **دوباره نمی‌نویسد**: کمک‌تابع‌های oxml موجود
پروژه از ``audit_report_generator.docx_builder`` (همان ماژولی که گزارش کمیسیون
کارگاه چک‌لیست را می‌سازد) به‌صورت فقط-خواندنی بازاستفاده می‌شوند:

    ``_set_paragraph_rtl``  -- علامت‌گذاری پاراگراف به‌عنوان bidi و راست‌چین
    ``_set_run_farsi``      -- تنظیم cs/ascii/hAnsi و زبان fa-IR روی run
    ``_set_table_rtl``      -- ``w:bidiVisual`` روی جدول
    ``_shade_cell``         -- سایه‌ی سلول (برای تأکید رنگی وضعیت)
    ``_mark_row_as_repeating_header`` -- تکرار سربرگ جدول در صفحات بعد
    ``_set_col_widths``     -- عرض ثابت ستون‌ها (جلوگیری از بیرون‌زدن جدول)

تنها تفاوت واقعی این کارگاه با آن ماژول، **فونت** است: پرامپت مرجع صراحتاً فونت
``Vazirmatn`` را می‌خواهد در حالی که ابزار موجود روی ``B Nazanin`` قفل شده است.
به همین دلیل این‌جا فونت یک پارامتر است و بعد از فراخوانی helper موجود، نام فونت
روی run بازنویسی می‌شود -- بدون کپی‌کردن منطق oxml و بدون تغییر فایل موجود.
"""
from __future__ import annotations

from typing import Callable, Optional, Sequence

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import _Cell

# بازاستفاده‌ی فقط-خواندنی از ابزار موجود پروژه (بدون هیچ تغییری در آن ماژول).
from audit_report_generator import docx_builder as _docx_builder

__all__ = [
    "DEFAULT_FONT",
    "PRIMARY_COLOR",
    "STATUS_SHADING",
    "add_bullet",
    "add_heading",
    "add_label_paragraph",
    "add_page_break",
    "add_paragraph",
    "add_table",
    "new_document",
    "status_shading",
]

# فونت اصلی گزارش (بخش «قالب‌بندی فایل خروجی» پرامپت مرجع). اگر روی سیستم بازکننده
# نصب نباشد، Word فونت جانشین انتخاب می‌کند -- مقدار همچنان در سند ثبت می‌ماند.
DEFAULT_FONT = "Vazirmatn"
PRIMARY_COLOR = "1F4E79"  # رنگ تیترها (هم‌خانواده‌ی رنگ گزارش‌های موجود پروژه)
MUTED_COLOR = "555555"
DANGER_COLOR = "B00000"

# رنگ‌بندی وضعیت: همیشه همراه با متن وضعیت استفاده می‌شود (هرگز تنها نشانگر نباشد).
STATUS_SHADING: dict[str, str] = {
    "مغایرت بااهمیت": "F8CBCB",
    "ریسک بااهمیت": "F8CBCB",
    "هشدار مدیریتی": "FCEBC5",
    "نزدیک به حد": "FCEBC5",
    "نیازمند بررسی": "FFF2CC",
    "فاقد داده کافی": "E8E8E8",
    "معاف": "EDEDED",
    "عادی": "DFF0D8",
}

_IMPORTANCE_SHADING: dict[str, str] = {
    "بسیار بالا": "F8CBCB",
    "بااهمیت": "FCEBC5",
    "قابل توجه": "FFF7E0",
}


def status_shading(status: str) -> Optional[str]:
    """رنگ پس‌زمینه‌ی متناظر با یک وضعیت (یا ``None`` اگر وضعیت ناشناخته باشد)."""
    return STATUS_SHADING.get(status)


def importance_shading(importance: str) -> Optional[str]:
    return _IMPORTANCE_SHADING.get(importance)


# ---------------------------------------------------------------------------
# پاراگراف و متن
# ---------------------------------------------------------------------------
def _apply_font(run, font: str) -> None:
    """نام فونت را روی همه‌ی دامنه‌های نویسه‌ای run بازنویسی می‌کند.

    helper موجود پروژه ``B Nazanin`` را روی ``w:cs/w:ascii/w:hAnsi`` می‌گذارد؛
    این‌جا فقط همان سه مقدار با فونت دلخواه جایگزین می‌شوند.
    """
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attribute in ("w:cs", "w:ascii", "w:hAnsi"):
        rFonts.set(qn(attribute), font)


def _set_run(
    run,
    *,
    size: float,
    bold: bool,
    color: Optional[str] = None,
    font: str = DEFAULT_FONT,
) -> None:
    # wiring زبان/RTL از ابزار موجود پروژه گرفته می‌شود، سپس فقط فونت بازنویسی می‌شود.
    _docx_builder._set_run_farsi(run, size=int(size), bold=bold)
    if font != _docx_builder.FARSI_FONT:
        _apply_font(run, font)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_paragraph(
    document,
    text: str,
    *,
    size: float = 11,
    bold: bool = False,
    color: Optional[str] = None,
    space_after: float = 6,
    space_before: float = 0,
    align_center: bool = False,
    font: str = DEFAULT_FONT,
    keep_with_next: bool = False,
):
    """یک پاراگراف فارسی راست‌چین (یا وسط‌چین) اضافه می‌کند."""
    paragraph = document.add_paragraph()
    _docx_builder._set_paragraph_rtl(paragraph, align_right=not align_center)
    if align_center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    run = paragraph.add_run(text)
    _set_run(run, size=size, bold=bold, color=color, font=font)
    if keep_with_next:
        _keep_with_next(paragraph)
    return paragraph


def add_label_paragraph(
    document,
    label: str,
    value: str,
    *,
    size: float = 11,
    space_after: float = 5,
    indent_cm: float = 0.0,
):
    """پاراگراف «برچسب: مقدار» با برچسب بولد (سبک بخش‌های ساختاریافته‌ی گزارش)."""
    paragraph = document.add_paragraph()
    _docx_builder._set_paragraph_rtl(paragraph)
    paragraph.paragraph_format.space_after = Pt(space_after)
    if indent_cm:
        paragraph.paragraph_format.right_indent = Cm(indent_cm)
    label_run = paragraph.add_run(f"{label}: ")
    _set_run(label_run, size=size, bold=True)
    value_run = paragraph.add_run(value)
    _set_run(value_run, size=size, bold=False)
    return paragraph


def add_heading(document, text: str, *, level: int = 1):
    """تیتر بخش: بولد، بزرگ‌تر با فاصله‌ی مناسب قبل/بعد و چسبیده به پاراگراف بعدی."""
    sizes = {1: 15, 2: 13, 3: 11.5}
    space_before = {1: 14, 2: 10, 3: 8}
    return add_paragraph(
        document,
        text,
        size=sizes.get(level, 11.5),
        bold=True,
        color=PRIMARY_COLOR,
        space_after=6,
        space_before=space_before.get(level, 8),
        keep_with_next=True,
    )


def add_bullet(document, text: str, *, size: float = 11, space_after: float = 3):
    paragraph = document.add_paragraph(style="List Bullet")
    _docx_builder._set_paragraph_rtl(paragraph)
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    _set_run(run, size=size, bold=False)
    return paragraph


def add_page_break(document) -> None:
    document.add_page_break()


# ---------------------------------------------------------------------------
# جدول
# ---------------------------------------------------------------------------
def _set_cell(
    cell: _Cell,
    text: str,
    *,
    size: float = 9.5,
    bold: bool = False,
    shading: Optional[str] = None,
    align_center: bool = False,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _docx_builder._set_paragraph_rtl(paragraph, align_right=not align_center)
    if align_center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.space_before = Pt(2)
    run = paragraph.add_run(text if text not in (None, "") else "—")
    _set_run(run, size=size, bold=bold)
    if shading:
        _docx_builder._shade_cell(cell, shading)
    cell.vertical_alignment = 1  # WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_row_cant_split(row) -> None:
    """جلوگیری از شکستن یک ردیف جدول بین دو صفحه (خروجی مرتب‌تر در Word)."""
    trPr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:cantSplit")
    element.set(qn("w:val"), "true")
    trPr.append(element)


def add_table(
    document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    *,
    widths_cm: Sequence[float],
    header_shading: str = "D9E2F3",
    cell_size: float = 9,
    header_size: float = 9.5,
    cant_split: bool = True,
    row_shading: Optional[Callable[[int, Sequence[str]], Optional[str]]] = None,
    bold_columns: Sequence[int] = (),
):
    """جدول راست‌به‌چپ با سربرگ تکرارشونده، عرض ثابت و متن Wrap‌شده.

    عرض کل ستون‌ها باید در عرض قابل‌استفاده‌ی صفحه (۲۱ سانتی‌متر منهای حاشیه‌ها)
    بماند؛ ``_set_col_widths`` همان ابزار موجود پروژه است که عرض را روی سلول‌ها و
    ستون‌ها ثابت می‌کند تا جدول از حاشیه بیرون نزند.
    """
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _docx_builder._set_table_rtl(table)

    header_row = table.rows[0]
    for index, header in enumerate(headers):
        _set_cell(header_row.cells[index], header, bold=True, size=header_size, shading=header_shading)
    _docx_builder._mark_row_as_repeating_header(header_row)
    if cant_split:
        _set_row_cant_split(header_row)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        if cant_split:
            _set_row_cant_split(row)
        shading = row_shading(row_index, values) if row_shading else None
        for column_index, value in enumerate(values):
            _set_cell(
                row.cells[column_index],
                str(value),
                size=cell_size,
                bold=column_index in bold_columns,
                shading=shading,
            )

    _docx_builder._set_col_widths(table, list(widths_cm))
    document.add_paragraph()
    return table


# ---------------------------------------------------------------------------
# سند
# ---------------------------------------------------------------------------
def new_document(*, font: str = DEFAULT_FONT, margin_cm: float = 2.0):
    """سند A4 با حاشیه‌ی متعارف و فونت پیش‌فرض فارسی."""
    from docx import Document

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(margin_cm)
    section.right_margin = Cm(margin_cm)
    section.top_margin = Cm(margin_cm)
    section.bottom_margin = Cm(margin_cm)

    # پیش‌فرض سبک Normal طوری تنظیم می‌شود که اگر جایی پاراگرافی بدون سبک صریح
    # ساخته شد، همچنان فارسی و خوانا بماند.
    normal = document.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(11)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attribute in ("w:cs", "w:ascii", "w:hAnsi"):
        rFonts.set(qn(attribute), font)
    return document
