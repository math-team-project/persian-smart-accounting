"""ساخت فایل خروجی Word (.docx) با فرمت‌بندی مناسب برای متن فارسی (راست‌به‌چپ).

از کتابخانه‌ی ``python-docx`` استفاده شده و تنظیمات RTL (بایدی/راست‌چین/فونت
اسکریپت پیچیده) به‌صورت دستی از طریق ``oxml`` اضافه می‌شوند، چون python-docx
به‌صورت پیش‌فرض این موارد را پشتیبانی نمی‌کند.
"""

from __future__ import annotations

import os
from datetime import datetime
from typing import List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import _Cell

from .schemas import AuditFinding, LoadedInput, ReportItem

FARSI_FONT = "B Nazanin"
FARSI_FALLBACK_FONT = "Tahoma"  # اگر B Nazanin روی سیستم بازکننده نصب نباشد

_SEVERITY_COLOR = {
    "بالا": "F8CBCB",     # قرمز کم‌رنگ
    "متوسط": "FCEBC5",    # زرد کم‌رنگ
    "کم": "DFF0D8",       # سبز کم‌رنگ
}
_SEVERITY_ORDER = {"بالا": 0, "متوسط": 1, "کم": 2}


# --------------------------------------------------------------------------
# کمک‌تابع‌های سطح پایین برای RTL / فونت
# --------------------------------------------------------------------------

def _set_paragraph_rtl(paragraph, align_right: bool = True) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    if align_right:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _set_run_farsi(run, size: Optional[int] = None, bold: Optional[bool] = None) -> None:
    run.font.name = FARSI_FONT
    run.font.complex_script = True
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:cs"), FARSI_FONT)
    rFonts.set(qn("w:ascii"), FARSI_FONT)
    rFonts.set(qn("w:hAnsi"), FARSI_FONT)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:bidi"), "fa-IR")
    rPr.append(lang)


def _add_farsi_paragraph(
    document,
    text: str,
    size: int = 12,
    bold: bool = False,
    align_right: bool = True,
    space_after: int = 6,
    color: Optional[str] = None,
) -> None:
    p = document.add_paragraph()
    _set_paragraph_rtl(p, align_right=align_right)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    _set_run_farsi(run, size=size, bold=bold)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_table_rtl(table) -> None:
    tblPr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tblPr.append(bidi)


def _set_cell_text(
    cell: _Cell,
    text: str,
    bold: bool = False,
    size: int = 10.5,
    align_right: bool = True,
    shading_hex: Optional[str] = None,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    _set_paragraph_rtl(p, align_right=align_right)
    run = p.add_run(text)
    _set_run_farsi(run, size=size, bold=bold)
    if shading_hex:
        _shade_cell(cell, shading_hex)
    cell.vertical_alignment = 1  # WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _shade_cell(cell: _Cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _mark_row_as_repeating_header(row) -> None:
    """این ردیف را به‌عنوان سربرگ تکرارشونده در صفحات بعدی جدول علامت می‌زند."""
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)


def _set_col_widths(table, widths_cm: List[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)
    for idx, width in enumerate(widths_cm):
        table.columns[idx].width = Cm(width)


# --------------------------------------------------------------------------
# ساخت گزارش
# --------------------------------------------------------------------------

def _format_number(value) -> str:
    if value is None:
        return "-"
    if isinstance(value, float):
        if value == int(value):
            return f"{int(value):,}"
        return f"{value:,.1f}"
    return str(value)


def _add_section_heading(document, text: str) -> None:
    _add_farsi_paragraph(document, text, size=13.5, bold=True, space_after=6, color="1F4E79")


def _add_checklist_section(document, report_items: List[ReportItem]) -> None:
    _add_section_heading(document, "بخش اول — موارد عدم تطابق شناسایی‌شده توسط چک‌لیست خودکار")
    _add_farsi_paragraph(
        document,
        "جدول زیر فقط موارد «عدم تطابق» را نشان می‌دهد که برای بررسی و تصمیم‌گیری کمیسیون "
        "آماده شده است. مواردی که علامت «بله» در ستون آخر دارند، پیش‌تر در گزارش رسمی "
        "حسابرس مستقل نیز به آن‌ها اشاره شده است.",
        size=10.5,
        color="555555",
        space_after=12,
    )

    if not report_items:
        _add_farsi_paragraph(
            document,
            "بر اساس نتایج چک‌لیست خودکار، در این دوره هیچ مورد عدم تطابقی شناسایی نشده است.",
            size=12,
            bold=True,
            space_after=14,
        )
        return

    # مرتب‌سازی بر اساس شدت (بالا -> کم) تا موارد مهم‌تر بالای جدول باشند
    ordered_items = sorted(report_items, key=lambda ri: _SEVERITY_ORDER.get(ri.severity, 1))

    headers = ["ردیف", "شناسه", "موضوع", "خلاصه مغایرت", "نکته کانونی جلسه", "اهمیت", "در گزارش حسابرسی؟"]
    col_widths = [1.2, 1.6, 3.3, 4.3, 4.3, 1.6, 2.2]

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    _set_table_rtl(table)

    header_row = table.rows[0]
    for idx, htext in enumerate(headers):
        _set_cell_text(header_row.cells[idx], htext, bold=True, size=10.5, shading_hex="D9D9D9")
    _mark_row_as_repeating_header(header_row)

    for i, ri in enumerate(ordered_items, start=1):
        row = table.add_row()
        values = [
            str(i),
            ri.question_id,
            ri.title,
            ri.summary,
            ri.committee_focus,
            ri.severity,
            "بله" if ri.mentioned_in_audit_report else "خیر",
        ]
        shading = _SEVERITY_COLOR.get(ri.severity)
        for idx, val in enumerate(values):
            # کل ردیف بر اساس شدت کمرنگ هایلایت می‌شود تا اهمیت مورد در نگاه اول دیده شود؛
            # ستون «موضوع» علاوه بر آن بولد می‌شود.
            _set_cell_text(
                row.cells[idx],
                val,
                size=10,
                bold=(idx == 2),
                shading_hex=shading,
            )

    _set_col_widths(table, col_widths)
    document.add_paragraph()


def _add_audit_findings_section(
    document, audit_findings: List[AuditFinding], error_message: Optional[str]
) -> None:
    _add_section_heading(
        document,
        "بخش دوم — سایر نکات مهم گزارش حسابرسی مستقل (خارج از پوشش چک‌لیست خودکار)",
    )
    _add_farsi_paragraph(
        document,
        "این بخش حاصل مرور کامل متن گزارش حسابرسی مستقل است و نکاتی را نشان می‌دهد که "
        "هیچ‌کدام از سوالات چک‌لیست خودکار به آن‌ها اشاره نکرده‌اند (مانند بند شرط یا "
        "تاکید حسابرس، ابهام در تداوم فعالیت، محدودیت در رسیدگی و موارد مشابه)؛ هدف از "
        "این بخش، پوشش کامل‌تر محتوای گزارش حسابرسی برای جلسه کمیسیون است.",
        size=10.5,
        color="555555",
        space_after=12,
    )

    if error_message:
        _add_farsi_paragraph(document, error_message, size=10.5, bold=True, color="B00000", space_after=14)
        return

    if not audit_findings:
        _add_farsi_paragraph(
            document,
            "در مرور متن گزارش حسابرسی، نکته‌ی مهم اضافه‌ای خارج از موارد بخش اول شناسایی نشد.",
            size=12,
            bold=True,
            space_after=14,
        )
        return

    ordered_findings = sorted(audit_findings, key=lambda af: _SEVERITY_ORDER.get(af.severity, 1))

    headers = ["ردیف", "موضوع", "خلاصه", "نکته کانونی جلسه", "اهمیت", "مرجع در گزارش"]
    col_widths = [1.2, 3.6, 4.8, 4.6, 1.6, 3.2]

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    _set_table_rtl(table)

    header_row = table.rows[0]
    for idx, htext in enumerate(headers):
        _set_cell_text(header_row.cells[idx], htext, bold=True, size=10.5, shading_hex="D9E2F3")
    _mark_row_as_repeating_header(header_row)

    for i, af in enumerate(ordered_findings, start=1):
        row = table.add_row()
        values = [
            str(i),
            af.title,
            af.summary,
            af.committee_focus,
            af.severity,
            af.source_hint or "-",
        ]
        shading = _SEVERITY_COLOR.get(af.severity)
        for idx, val in enumerate(values):
            _set_cell_text(
                row.cells[idx],
                val,
                size=10,
                bold=(idx == 1),
                shading_hex=shading,
            )

    _set_col_widths(table, col_widths)
    document.add_paragraph()


def build_docx(
    loaded: LoadedInput,
    report_items: List[ReportItem],
    audit_findings: List[AuditFinding],
    output_path: str,
    report_title: str,
    entity_name: Optional[str] = None,
    audit_findings_error: Optional[str] = None,
    has_audit_report: bool = True,
) -> str:
    """فایل docx نهایی را می‌سازد و در ``output_path`` ذخیره می‌کند.

    اگر ``has_audit_report=False`` باشد (یعنی کاربر گزارش حسابرسی نداده)، فایل خروجی
    فقط شامل بخش اول (خلاصه‌ی موارد چک‌لیست) خواهد بود و بخش دوم به‌طور کامل حذف
    می‌شود — نه این‌که به‌عنوان خطا نمایش داده شود.

    Returns
    -------
    str
        مسیر فایل ذخیره‌شده (همان output_path).
    """
    document = Document()

    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    # --- عنوان ---
    _add_farsi_paragraph(document, report_title, size=16, bold=True, space_after=4)

    meta_bits = []
    if entity_name:
        meta_bits.append(f"سازمان: {entity_name}")
    meta_bits.append(f"تاریخ تهیه گزارش: {datetime.now().strftime('%Y-%m-%d')}")
    if has_audit_report:
        meta_bits.append("موضوع: جمع‌بندی چک‌لیست حسابرسی و گزارش حسابرسی مستقل جهت طرح در جلسه کمیسیون")
    else:
        meta_bits.append("موضوع: خلاصه موارد عدم تطابق چک‌لیست حسابرسی جهت طرح در جلسه کمیسیون")
    _add_farsi_paragraph(document, " | ".join(meta_bits), size=10.5, color="555555", space_after=14)

    if not has_audit_report:
        _add_farsi_paragraph(
            document,
            "توجه: گزارش حسابرسی مستقل برای این اجرا ارائه نشده است؛ این گزارش صرفاً بر "
            "اساس نتایج چک‌لیست خودکار تهیه شده و شامل تحلیل متن گزارش حسابرسی نیست.",
            size=10,
            color="777777",
            space_after=12,
        )

    # --- خلاصه وضعیت کلی چک‌لیست ---
    if loaded.all_items_count is not None:
        rate = loaded.compliance_rate
        rate_txt = f"{rate:.1f}٪" if isinstance(rate, (int, float)) else "نامشخص"
        summary_text = (
            f"از مجموع {_format_number(loaded.all_items_count)} سوال بررسی‌شده در چک‌لیست حسابرسی، "
            f"{_format_number(loaded.true_count)} مورد منطبق، "
            f"{_format_number(loaded.false_count)} مورد دارای عدم تطابق"
        )
        extra = []
        if loaded.error_count:
            extra.append(f"{_format_number(loaded.error_count)} مورد خطای فنی در ارزیابی خودکار")
        if loaded.manual_count:
            extra.append(f"{_format_number(loaded.manual_count)} مورد نیازمند بررسی دستی حسابرس")
        if extra:
            summary_text += "، " + "، ".join(extra)
        summary_text += f" بوده است. نرخ تطابق کلی: {rate_txt}."
        _add_farsi_paragraph(document, summary_text, size=11.5, space_after=14)

    # --- بخش اول: چک‌لیست ---
    _add_checklist_section(document, report_items)

    # --- بخش دوم: نکات مهم گزارش حسابرسی خارج از چک‌لیست (فقط اگر گزارش حسابرسی داده شده) ---
    if has_audit_report:
        _add_audit_findings_section(document, audit_findings, audit_findings_error)

    _add_farsi_paragraph(
        document,
        "این گزارش صرفاً جهت جمع‌بندی و تسهیل تصمیم‌گیری در جلسه کمیسیون تهیه شده و "
        "جایگزین گزارش رسمی حسابرس مستقل نیست؛ برای جزئیات فنی و مبنای محاسباتی هر مورد "
        "باید به گزارش حسابرسی و یادداشت‌های توضیحی صورت‌های مالی مراجعه شود.",
        size=9.5,
        color="777777",
        space_after=0,
    )

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    document.save(output_path)
    return output_path
